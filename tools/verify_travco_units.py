#!/usr/bin/env python3
"""Check the Travco units on the site against the client price sheet.

The seventeen Makadina and Marina Gate units were written from the sheet, and
the thing most easily got wrong is the images: a unit may name several, the
order is the sheet's, and two units may legitimately name the same file in a
different order. This reads the workbook again, independently of whatever
generated the site, and asserts for each unit in sheet order:

  price, area, bedrooms, bathrooms, down payment, instalment years, availability
  the render list, in the sheet's order, first image first
  the floor plans, the master plan and the location map
  that no unit shows an image the sheet does not name for it

    python3 tools/verify_travco_units.py [path/to/the.xlsx]
"""
import os
import re
import sys

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DEFAULT_XLSX = ('/root/.claude/uploads/c9c8c82a-00eb-5d87-a89f-ba9e63e19221/'
                'dedaa386-________________.xlsx')

PROJ = {'MAKADINA': 'makadina', 'MARINAGATE': 'marina-gate'}
IDS = ['MK-%02d' % i for i in range(1, 11)] + ['MG-%02d' % i for i in range(1, 8)]


def slug(s):
    return re.sub(r'[^a-z0-9]+', '-', s.lower()).strip('-')


def sheet_units(path):
    """The TRAVCO rows and the image ids listed under each, in sheet order."""
    ws = openpyxl.load_workbook(path, data_only=True)['الوحدات']
    out, cur = [], None
    for i in range(3, ws.max_row + 1):
        r = [c.value for c in ws[i]][:15]
        dev = str(r[0]).strip().upper() if r[0] else ''
        if dev.startswith('TRAVCO'):
            cur = dict(row=i, proj=str(r[1]).strip().upper(), bed=r[4], sqm=r[5],
                       price=int(str(r[6]).replace(',', '').strip()),
                       dp=r[7], yrs=r[8], avail=r[9], bath=r[14],
                       imgs=[str(r[10]).strip()] if r[10] else [],
                       fps=[str(r[12]).strip()] if r[12] else [],
                       mp=str(r[11]).strip() if r[11] else None,
                       loc=str(r[13]).strip() if r[13] else None)
            out.append(cur)
        elif dev:
            cur = None
        elif cur and all(v in (None, '') for v in r[:10]):
            if r[10]:
                cur['imgs'].append(str(r[10]).strip())
            if r[12]:
                cur['fps'].append(str(r[12]).strip())
    return out


def main():
    xlsx = sys.argv[1] if len(sys.argv) > 1 else DEFAULT_XLSX
    if not os.path.exists(xlsx):
        sys.exit('the sheet is not at %s — pass its path as an argument' % xlsx)

    us = sheet_units(xlsx)
    recs = open(os.path.join(ROOT, 'src/tpl_script2.html'), encoding='utf-8').read()
    media = open(os.path.join(ROOT, 'src/tpl_script2b.html'), encoding='utf-8').read()
    bad = 0

    def check(cond, msg):
        nonlocal bad
        if not cond:
            print('   FAIL  %s' % msg)
            bad += 1

    def listed(var, uid):
        m = re.search(r"^    '%s':\[(.*?)\],$" % uid, media, re.M)
        if not m:
            m2 = re.search(r"^    '%s':'([^']+)',$" % uid, media, re.M)
            return [m2.group(1)] if m2 else None
        return re.findall(r"'([^']+)'", m.group(1))

    def block(var, uid):
        seg = media[media.index('var %s = {' % var):]
        seg = seg[:seg.index('\n  };')]
        m = re.search(r"^    '%s':(\[(.*?)\]|'[^']*'),$" % uid, seg, re.M)
        return re.findall(r"'(/[^']+)'", m.group(0)) if m else None

    check(len(us) == len(IDS), 'the sheet has %d Travco rows, the site %d' % (len(us), len(IDS)))
    print('%-7s %-13s %-5s %-6s %s' % ('unit', 'price', 'm2', 'images', 'first image'))
    for uid, u in zip(IDS, us):
        line = [l for l in recs.split('\n') if "id:'%s'" % uid in l]
        check(bool(line), '%s: no record on the site' % uid)
        if not line:
            continue
        line = line[0]
        for field, val in (('beds', u['bed']), ('baths', u['bath']),
                           ('area', u['sqm']), ('price', u['price']),
                           ('dp', round(u['dp'] * 100)), ('years', u['yrs'])):
            check('%s:%s,' % (field, val) in line or '%s:%s}' % (field, val) in line,
                  '%s: %s should be %s' % (uid, field, val))
        check("handover:'%s'" % u['avail'] in line, '%s: availability %s' % (uid, u['avail']))

        base = '/project-media/travco/%s/units/' % PROJ[u['proj']]
        want_g = [base + slug(n) + '.webp' for n in u['imgs']]
        got_g = block('UNIT_GALLERY', uid)
        check(got_g == want_g,
              '%s: gallery out of order\n         sheet %s\n         site  %s' % (uid, want_g, got_g))
        check(block('UNIT_IMAGES', uid) == want_g[:1], '%s: cover is not the sheet\'s first image' % uid)
        check(block('UNIT_FLOORPLANS', uid) == [base + slug(n) + '.webp' for n in u['fps']],
              '%s: floor plans' % uid)
        check(block('UNIT_MASTERPLANS', uid) == [base + slug(u['mp']) + '.webp'], '%s: master plan' % uid)
        check(block('UNIT_LOCATIONS', uid) == [base + slug(u['loc']) + '.webp'], '%s: location' % uid)

        for path in (got_g or []):
            check(os.path.exists(os.path.join(ROOT, path.lstrip('/'))),
                  '%s: %s is referenced but not in the repo' % (uid, path))

        print('%-7s %-13s %-5s %-6d %s'
              % (uid, '{:,}'.format(u['price']), u['sqm'], len(u['imgs']), u['imgs'][0]))

    print('\n%d checks failed' % bad)
    return 1 if bad else 0


if __name__ == '__main__':
    sys.exit(main())
