#!/usr/bin/env python3
"""Check the reloaded SODIC units on the site against the client price sheet.

Six SODIC projects are written from the sheet, and the thing most easily got
wrong is the images: a unit may name several, the order is the sheet's, and a
continuation row carries nothing but another image id. This reads the workbook
again, independently of whatever generated the site, and asserts for each unit
in sheet order:

  price, area, bedrooms, bathrooms, down payment, instalment years, handover
  the render list, in the sheet's order, first image first
  the floor plans, the master plan and the location map
  that every path it shows is a file that exists in the repo

An id the sheet names that no archive supplies is reported at the end rather
than failed: the unit shows what it has rather than a broken picture, and the
report says exactly which file is still wanted.

    python3 tools/verify_sodic_units.py [path/to/the.xlsx]
"""
import json
import os
import re
import subprocess
import sys
import tempfile

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DEFAULT_XLSX = ('/root/.claude/uploads/c9c8c82a-00eb-5d87-a89f-ba9e63e19221/'
                'dedaa386-________________.xlsx')

# the sheet's project wording -> the site's slug and unit-id prefix, in the
# order the site lists them. Note the sheet spells The Estates and The Estates
# Residence as two separate groups, and they are two separate projects here.
SPEC = [('Villette', 'villette', 'VL'),
        ('june', 'june-north-coast', 'JN'),
        ('SODIC EAST', 'sodic-east', 'SE'),
        ('THE ESTATES', 'the-estates-zayed', 'ES'),
        ('Westown Medical Center', 'westown-medical-center', 'WM'),
        ('The Estates Residence', 'the-estates-residence', 'TR')]


def slug(s):
    return re.sub(r'[^a-z0-9]+', '-', s.strip().lower()).strip('-')


def sheet_units(path):
    """The SODIC rows, grouped by the sheet's project name, in sheet order.

    A unit is one data row plus the continuation rows beneath it, which carry
    nothing but a further image id (column 10) and sometimes a further floor
    plan (column 12). That is how a unit comes to have more than one picture.
    """
    ws = openpyxl.load_workbook(path, data_only=True)['الوحدات']
    out, cur = {}, None
    for i in range(3, ws.max_row + 1):
        r = [c.value for c in ws[i]][:15]
        dev = str(r[0]).strip().lower() if r[0] else ''
        if dev == 'sodic':
            cur = dict(row=i, proj=str(r[1]).strip(), type=str(r[3]).strip(),
                       bed=r[4], sqm=r[5],
                       price=int(str(r[6]).replace(',', '').strip()),
                       dp=r[7], yrs=r[8], avail=r[9], bath=r[14],
                       imgs=[str(r[10]).strip()] if r[10] else [],
                       fps=[str(r[12]).strip()] if r[12] else [],
                       mp=str(r[11]).strip() if r[11] else None,
                       loc=str(r[13]).strip() if r[13] else None)
            out.setdefault(cur['proj'], []).append(cur)
        elif dev:
            cur = None
        elif cur and all(v in (None, '') for v in r[:10]):
            if r[10]:
                cur['imgs'].append(str(r[10]).strip())
            if r[12]:
                cur['fps'].append(str(r[12]).strip())
    return out


def site_units():
    """The unit records as the browser evaluates them, keyed by code.

    Reading the source text is not enough. The sheet writes a size band as
    "153-157", and a generator that dropped that into the record verbatim
    produced `area:153-157` — valid JavaScript that evaluates to -4, while the
    source still reads exactly like the sheet. Comparing against the running
    bundle is the only way that class of mistake shows up.
    """
    # Through a file, not a pipe: the dump is several hundred KB and a captured
    # pipe truncates it, which reads as a parse error rather than a short read.
    with tempfile.NamedTemporaryFile('w+', suffix='.json', delete=False) as fh:
        tmp = fh.name
    try:
        with open(tmp, 'w') as fh:
            subprocess.run(['node', os.path.join(ROOT, 'tools/domtest.cjs'), '--dump-data'],
                           stdout=fh, stderr=subprocess.DEVNULL, cwd=ROOT, check=True)
        with open(tmp, encoding='utf-8') as fh:
            return {u['id']: u for u in json.load(fh)['UNITS']}
    finally:
        os.unlink(tmp)


def main():
    xlsx = sys.argv[1] if len(sys.argv) > 1 else DEFAULT_XLSX
    if not os.path.exists(xlsx):
        sys.exit('the sheet is not at %s — pass its path as an argument' % xlsx)

    groups = sheet_units(xlsx)
    live = site_units()
    media = open(os.path.join(ROOT, 'src/tpl_script2b.html'), encoding='utf-8').read()
    bad, absent = 0, []

    def check(cond, msg):
        nonlocal bad
        if not cond:
            print('   FAIL  %s' % msg)
            bad += 1

    def block(var, uid):
        """The paths one map holds for a unit, or None when it holds no row."""
        seg = media[media.index('var %s = {' % var):]
        seg = seg[:seg.index('\n  };')]
        m = re.search(r"^    '%s':(\[(.*?)\]|'[^']*'),$" % uid, seg, re.M)
        return re.findall(r"'(/[^']+)'", m.group(0)) if m else None

    print('%-7s %-13s %-7s %-6s %s' % ('unit', 'price', 'm2', 'images', 'first image'))
    for name, slug_, code in SPEC:
        rows = groups.get(name) or []
        check(bool(rows), '%s: the sheet has no rows under this name' % name)
        base = '/project-media/sodic/%s/units/' % slug_

        def path(n):
            """The stored path for an image id, or None when no file exists.

            The Estates Residence names its master plan "mp the eastates." —
            a trailing dot the file does not carry — so the id is slugified,
            which drops it, exactly as the loader does.
            """
            p = base + slug(n) + '.webp'
            if os.path.exists(os.path.join(ROOT, p.lstrip('/'))):
                return p
            absent.append((name, n))
            return None

        for i, u in enumerate(rows, 1):
            uid = '%s-%02d' % (code, i)
            rec = live.get(uid)
            check(bool(rec), '%s: no record on the site (sheet row %d)' % (uid, u['row']))
            if not rec:
                continue
            check(rec['project'] == slug_, '%s: is not on %s' % (uid, slug_))

            # A size the sheet gives as a band ("153-157") is two fields on the
            # site, so compare the pair against the band's own two ends.
            sqm = str(u['sqm']).strip()
            lo, _, hi = sqm.partition('-') if '-' in sqm else (sqm, '', '')
            want = [('price', u['price']), ('dp', round(u['dp'] * 100)),
                    ('years', u['yrs']), ('area', float(lo))] \
                + ([('areaTo', float(hi))] if hi else []) \
                + [(f, v) for f, v in (('beds', u['bed']), ('baths', u['bath']))
                   if v not in (None, '')]
            # A clinic carries neither count: the sheet leaves both blank and an
            # invented 0 would render on the page as one.
            for f, v in (('beds', u['bed']), ('baths', u['bath'])):
                check(v not in (None, '') or rec.get(f) is None,
                      '%s: the sheet gives no %s, the site shows %s' % (uid, f, rec.get(f)))
            for field, val in want:
                got = rec.get(field)
                check(got is not None and float(got) == float(val),
                      '%s: %s is %s, the sheet says %s (row %d)'
                      % (uid, field, got, val, u['row']))
            check(not hi or 'areaTo' in rec, '%s: the sheet gives a size band, the site one number' % uid)
            check(str(rec.get('handover')) == str(u['avail']),
                  '%s: handover %s, the sheet says %s' % (uid, rec.get('handover'), u['avail']))

            # Images the archives supply, in the sheet's order. An id with no
            # file is dropped rather than shown broken, so the site's list is
            # the sheet's list minus whatever was never delivered.
            want_g = [p for p in (path(n) for n in u['imgs']) if p]
            got_g = block('UNIT_GALLERY', uid) or []
            check(got_g == want_g,
                  '%s: gallery out of order\n         sheet %s\n         site  %s'
                  % (uid, want_g, got_g))
            check((block('UNIT_IMAGES', uid) or []) == want_g[:1],
                  '%s: cover is not the sheet\'s first image' % uid)
            check((block('UNIT_FLOORPLANS', uid) or []) ==
                  [p for p in (path(n) for n in u['fps']) if p], '%s: floor plans' % uid)
            check((block('UNIT_MASTERPLANS', uid) or []) ==
                  [p for p in [path(u['mp'])] if p] if u['mp'] else True, '%s: master plan' % uid)
            check((block('UNIT_LOCATIONS', uid) or []) ==
                  [p for p in [path(u['loc'])] if p] if u['loc'] else True, '%s: location' % uid)

            print('%-7s %-13s %-7s %-6s %s'
                  % (uid, '{:,}'.format(u['price']), u['sqm'], '%d/%d' % (len(want_g), len(u['imgs'])),
                     u['imgs'][0] if u['imgs'] else '—'))

    print('\n%d checks failed' % bad)
    seen = sorted({(p, n) for p, n in absent})
    if seen:
        print('%d image ids the sheet names that no archive supplies:' % len(seen))
        for p, n in seen:
            print('   %-12s %s' % (p, n))
    return 1 if bad else 0


if __name__ == '__main__':
    sys.exit(main())
